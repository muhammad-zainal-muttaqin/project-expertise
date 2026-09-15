"""Ubah laporan Markdown menjadi DOCX memakai python-docx, tanpa pandoc.

Konstruksi yang ditangani mengikuti kebutuhan laporan monev: judul, paragraf,
daftar berbutir dan bernomor bersarang, tabel dengan pipa ber-escape, gambar,
keterangan gambar miring, serta format sebaris tebal, miring, kode, dan tautan.
Tautan dipertahankan sebagai pranala aktif.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


FONT_UTAMA, FONT_KODE = "Times New Roman", "Consolas"
TEAL = RGBColor(0x17, 0x6B, 0x72)
LEBAR_ISI_CM = 17.0
BASE_SUMBER = Path.cwd()
INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`|\*[^*\s][^*]*?\*)")
PIPA = re.compile(r"(?<!\\)\|")
DAFTAR = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
GAMBAR = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")


def atur_font(element, nama):
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), nama)
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)


def bersih(teks):
    teks = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", teks)
    return teks.replace("**", "").replace("`", "").replace("\\|", "|")


def tambah_inline(par, teks, ukuran=None, tebal=False, miring=False):
    for bagian in INLINE.split(teks):
        if not bagian:
            continue
        if bagian.startswith("**") and bagian.endswith("**") and len(bagian) > 4:
            tambah_inline(par, bagian[2:-2], ukuran, True, miring)
        elif bagian.startswith("[") and "](" in bagian:
            label, tujuan = bagian[1:].split("](", 1)
            tujuan = tujuan[:-1]
            if not tujuan.startswith(("http:", "https:", "#")):
                jalur, _, fragmen = tujuan.partition("#")
                tujuan = (BASE_SUMBER / jalur).resolve().as_uri()
                if fragmen:
                    tujuan += "#" + fragmen
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), par.part.relate_to(tujuan, RT.HYPERLINK, is_external=True))
            run = par.add_run(bersih(label).replace("*", ""))
            run.bold, run.italic = tebal, miring
            run.font.size = Pt(ukuran or 11)
            run.font.color.rgb = RGBColor(0x17, 0x6B, 0x72)
            link.append(run._r)
            par._p.append(link)
        elif bagian.startswith("`") and bagian.endswith("`") and len(bagian) > 2:
            run = par.add_run(bagian[1:-1].replace("\\|", "|"))
            run.bold, run.italic = tebal, miring
            atur_font(run._r, FONT_KODE)
            run.font.size = Pt((ukuran or 11) - 1.5)
        elif bagian.startswith("*") and bagian.endswith("*") and len(bagian) > 2:
            tambah_inline(par, bagian[1:-1], ukuran, tebal, True)
        else:
            run = par.add_run(bagian.replace("\\|", "|"))
            run.bold, run.italic = tebal, miring
            if ukuran:
                run.font.size = Pt(ukuran)


def pisah_sel(baris):
    isi = baris.strip()
    if isi.startswith("|"):
        isi = isi[1:]
    if isi.endswith("|") and not isi.endswith("\\|"):
        isi = isi[:-1]
    return [sel.strip() for sel in PIPA.split(isi)]


def lebar_kolom(judul, data, kanan, ukuran):
    """Bobot kolom menurut panjang isi, dengan batas bawah selebar kata terpanjang.

    Batas bawah mencegah Word memotong angka atau kata di tengah; perkiraan
    lebar satu karakter 0,02 cm per poin ditambah 0,5 cm margin sel.
    """
    bobot, minimum = [], []
    for c, kepala in enumerate(judul):
        sel = [bersih(b[c]) for b in data if c < len(b)]
        isi = sorted(len(s) for s in sel) or [0]
        kata = max((len(w) for s in sel + [bersih(kepala)] for w in s.split()), default=0)
        if kanan[c]:
            nilai = max(len(bersih(kepala)), isi[-1], 6)
        else:
            nilai = max(isi[int(.8 * (len(isi) - 1))], len(bersih(kepala)), 4)
        bobot.append(min(nilai, 90) ** .85)
        minimum.append(min(kata * ukuran * .02 + .5, .45 * LEBAR_ISI_CM))
    tetap = set()
    while True:
        bebas = [c for c in range(len(judul)) if c not in tetap]
        sisa = LEBAR_ISI_CM - sum(minimum[c] for c in tetap)
        if not bebas or sisa <= 0:
            skala = LEBAR_ISI_CM / sum(minimum)
            return [m * skala for m in minimum]
        total = sum(bobot[c] for c in bebas)
        lebar = [minimum[c] if c in tetap else sisa * bobot[c] / total for c in range(len(judul))]
        kurang = [c for c in bebas if lebar[c] < minimum[c]]
        if not kurang:
            return lebar
        tetap.update(kurang)


def arsir(sel, warna):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), warna)
    sel._tc.get_or_add_tcPr().append(shd)


def tambah_tabel(doc, baris):
    judul = pisah_sel(baris[0])
    kanan = [s.endswith(":") and not s.startswith(":") for s in pisah_sel(baris[1])]
    data = [pisah_sel(b) for b in baris[2:]]
    n = len(judul)
    ukuran = 10 if n <= 4 else 9.5 if n <= 6 else 9
    lebar = lebar_kolom(judul, data, kanan, ukuran)
    tabel = doc.add_table(rows=1 + len(data), cols=n)
    tabel.style = "Table Grid"
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabel.autofit = False
    properti = tabel._tbl.tblPr
    batas = OxmlElement("w:tblBorders")
    for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
        garis = OxmlElement(f"w:{sisi}")
        for nama, nilai in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            garis.set(qn(f"w:{nama}"), nilai)
        batas.append(garis)
    properti.append(batas)
    margin = OxmlElement("w:tblCellMar")
    for sisi, nilai in (("top", "50"), ("bottom", "50"), ("left", "85"), ("right", "85")):
        tepi = OxmlElement(f"w:{sisi}")
        tepi.set(qn("w:w"), nilai)
        tepi.set(qn("w:type"), "dxa")
        margin.append(tepi)
    properti.append(margin)
    for c in range(n):
        tabel.columns[c].width = Cm(lebar[c])
    for r, isi_baris in enumerate([judul] + data):
        tabel.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c in range(n):
            sel = tabel.rows[r].cells[c]
            sel.width = Cm(lebar[c])
            sel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par = sel.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.05
            if kanan[c]:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            tambah_inline(par, isi_baris[c] if c < len(isi_baris) else "", ukuran, tebal=r == 0)
            if r == 0:
                arsir(sel, "E7ECEF")
                par.paragraph_format.keep_with_next = True
    kepala = tabel.rows[0]._tr.get_or_add_trPr()
    ulang = OxmlElement("w:tblHeader")
    ulang.set(qn("w:val"), "true")
    kepala.append(ulang)
    sela = doc.add_paragraph()
    sela.paragraph_format.space_after = Pt(4)
    sela.paragraph_format.space_before = Pt(0)
    sela.paragraph_format.line_spacing = Pt(2)
    sela.add_run().font.size = Pt(2)


def nomor_halaman(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for jenis, instruksi in (("begin", None), (None, "PAGE"), ("end", None)):
        if jenis:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), jenis)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruksi
        run._r.append(el)


def siapkan_dokumen():
    doc = Document()
    for batas in list(doc.styles.element.iter(qn("w:pBdr"))):
        batas.getparent().remove(batas)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for tepi in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, tepi, Cm(2))
    normal = doc.styles["Normal"]
    atur_font(normal.element, FONT_UTAMA)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for nama, ukuran, warna in (("Title", 17, RGBColor(0, 0, 0)), ("Heading 1", 14, RGBColor(0, 0, 0)), ("Heading 2", 12, RGBColor(0, 0, 0))):
        gaya = doc.styles[nama]
        atur_font(gaya.element, FONT_UTAMA)
        gaya.font.size = Pt(ukuran)
        gaya.font.bold = True
        gaya.font.color.rgb = warna
        gaya.paragraph_format.space_before = Pt(0 if nama == "Title" else 14)
        gaya.paragraph_format.space_after = Pt(6)
        gaya.paragraph_format.keep_with_next = True
    nomor_halaman(sec)
    return doc


def ubah(sumber, keluaran):
    global BASE_SUMBER
    BASE_SUMBER = sumber.resolve().parent
    doc = siapkan_dokumen()
    baris = sumber.read_text(encoding="utf-8").splitlines()
    paragraf = []

    def tutup_paragraf():
        if paragraf:
            tambah_inline(doc.add_paragraph(), " ".join(paragraf))
            paragraf.clear()

    i = 0
    while i < len(baris):
        teks = baris[i].rstrip()
        daftar = DAFTAR.match(teks)
        gambar = GAMBAR.match(teks)
        if not teks.strip() or teks.strip() == "---":
            tutup_paragraf()
        elif teks.startswith("#"):
            tutup_paragraf()
            level = len(teks) - len(teks.lstrip("#"))
            judul = bersih(teks[level:].strip()).replace("*", "")
            if level == 1:
                doc.add_paragraph(judul, style="Title")
            else:
                heading = doc.add_heading(judul, level=min(level - 1, 2))
                if judul.startswith("Lampiran C"):
                    heading.paragraph_format.page_break_before = True
        elif teks.startswith("|"):
            tutup_paragraf()
            blok = []
            while i < len(baris) and baris[i].startswith("|"):
                blok.append(baris[i])
                i += 1
            tambah_tabel(doc, blok)
            continue
        elif gambar:
            tutup_paragraf()
            doc.add_picture(str(sumber.parent / gambar.group(2)), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        elif teks.startswith("*") and teks.endswith("*") and not teks.startswith("**"):
            tutup_paragraf()
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tambah_inline(par, teks[1:-1], 9.5, miring=True)
        elif teks.startswith("```"):
            tutup_paragraf()
            i += 1
            while i < len(baris) and not baris[i].startswith("```"):
                run = doc.add_paragraph().add_run(baris[i])
                atur_font(run._r, FONT_KODE)
                run.font.size = Pt(9)
                i += 1
        elif daftar:
            tutup_paragraf()
            level = len(daftar.group(1).replace("\t", "   ")) // 3
            penanda = daftar.group(2)
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(.8 + .7 * level)
            par.paragraph_format.first_line_indent = Cm(-.6)
            par.paragraph_format.space_after = Pt(3)
            par.add_run(("•" if penanda in "-*" else penanda) + "\t")
            par.paragraph_format.tab_stops.add_tab_stop(Cm(.8 + .7 * level))
            tambah_inline(par, daftar.group(3))
        else:
            paragraf.append(teks.strip())
        i += 1
    tutup_paragraf()
    doc.save(keluaran)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sumber", type=Path)
    parser.add_argument("-o", "--keluaran", type=Path)
    args = parser.parse_args()
    keluaran = args.keluaran or args.sumber.with_suffix(".docx")
    ubah(args.sumber, keluaran)
    print(keluaran)


if __name__ == "__main__":
    main()
